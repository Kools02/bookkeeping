from __future__ import annotations
from io import BytesIO

import io
import os
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Optional

from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func, or_

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
EXPORT_DIR = BASE_DIR / "exports"
DB_PATH = BASE_DIR / "tin_app.db"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "change-this-secret")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", f"sqlite:///{DB_PATH}")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB

db = SQLAlchemy(app)


class ReferenceEntry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    actual_tin_receipt = db.Column(db.String(64), nullable=False)
    actual_tin_normalized = db.Column(db.String(32), nullable=False, index=True)
    tin_needed_input_tax = db.Column(db.String(64), nullable=False)
    store = db.Column(db.String(255), nullable=False)
    address = db.Column(db.String(255), nullable=False)
    category = db.Column(db.String(64), nullable=False)
    uploaded_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)


class Period(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(32), nullable=False, unique=True)
    is_active = db.Column(db.Boolean, nullable=False, default=False)
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)


class Entry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    period_id = db.Column(db.Integer, db.ForeignKey("period.id"), nullable=False, index=True)
    category = db.Column(db.String(64), nullable=True)
    store = db.Column(db.String(255), nullable=True)
    address = db.Column(db.String(255), nullable=True)
    actual_tin_number = db.Column(db.String(64), nullable=False)
    tin_number_on_input_tax = db.Column(db.String(64), nullable=True)
    invoice_value = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    vat_12 = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    cost_wo_vat = db.Column(db.Numeric(12, 2), nullable=False, default=0)
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow, onupdate=datetime.utcnow)

    period = db.relationship("Period", backref=db.backref("entries", lazy=True, cascade="all, delete-orphan"))


HEADER_ALIASES = {
    "actual_tin_number_on_receipt": "actual_tin_number_on_receipt",
    "actualtinnumberonreceipt": "actual_tin_number_on_receipt",
    "tin_number_needed_on_input_tax": "tin_number_needed_on_input_tax",
    "tinnumberneededoninputtax": "tin_number_needed_on_input_tax",
    "tin_number_on_input_tax": "tin_number_needed_on_input_tax",
    "store": "store",
    "address": "address",
    "cat": "cat",
    "category": "cat",
}

REQUIRED_HEADER_KEYS = {
    "actual_tin_number_on_receipt",
    "tin_number_needed_on_input_tax",
    "store",
    "address",
    "cat",
}

TWOPLACES = Decimal("0.01")

CATEGORY_CANONICAL_MAP = {
    "DSGROCERY": "D/S-GROCERY",
    "DSSTORE": "D/S-STORE",
    "ELECWATER": "ELEC/WATER",
    "ELECTWATER": "ELEC/WATER",
    "GO": "G&O",
    "HM": "H/M",
    "MED": "MED",
    "OS": "O/S",
    "RENT": "RENT",
    "REP": "REP",
    "REPSPECIAL": "REP-SPECIAL",
    "RM": "R/M",
    "TEL": "TEL",
    "TP": "T/P",
    "TRANSPO": "TRANSPO",
}


def normalize_tin(value: str | None) -> str:
    if value is None:
        return ""
    return re.sub(r"\D", "", str(value).strip())


def canonical_header(value: object) -> str:
    raw = str(value or "").strip().lower()
    raw = re.sub(r"[^a-z0-9]+", "_", raw).strip("_")
    return HEADER_ALIASES.get(raw, raw)


def normalize_category(value: str | None) -> str:
    raw = str(value or "").strip().upper()
    if not raw:
        return ""
    raw = raw.replace("—", "-").replace("–", "-")
    raw = re.sub(r"\s+", " ", raw).strip()
    key = re.sub(r"[^A-Z0-9]", "", raw)
    return CATEGORY_CANONICAL_MAP.get(key, raw)


def to_decimal(raw: str | float | int | Decimal | None) -> Decimal:
    if raw is None or raw == "":
        return Decimal("0.00")
    if isinstance(raw, Decimal):
        return raw.quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    text = str(raw).replace(",", "").strip()
    try:
        return Decimal(text).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        raise ValueError("Invoice value must be a valid number.")


def compute_vat_and_cost(invoice_value: Decimal) -> tuple[Decimal, Decimal]:
    vat = (invoice_value / Decimal("1.12") * Decimal("0.12")).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    cost = (invoice_value - vat).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    return vat, cost


def get_reference_categories() -> list[str]:
    rows = (
        db.session.query(ReferenceEntry.category)
        .filter(ReferenceEntry.category.isnot(None))
        .filter(func.trim(ReferenceEntry.category) != "")
        .all()
    )
    return sorted({normalize_category(row[0]) for row in rows if normalize_category(row[0])})


def cleanup_existing_categories() -> None:
    changed = False

    for row in ReferenceEntry.query.all():
        canonical = normalize_category(row.category)
        if canonical and row.category != canonical:
            row.category = canonical
            changed = True

    for entry in Entry.query.all():
        canonical = normalize_category(entry.category)
        if canonical and entry.category != canonical:
            entry.category = canonical
            changed = True

    seen_reference_keys: set[tuple[str, str, str, str, str]] = set()
    for row in ReferenceEntry.query.order_by(ReferenceEntry.id.asc()).all():
        key = (
            row.actual_tin_normalized or "",
            (row.tin_needed_input_tax or "").strip(),
            (row.store or "").strip(),
            (row.address or "").strip(),
            (row.category or "").strip(),
        )
        if key in seen_reference_keys:
            db.session.delete(row)
            changed = True
            continue
        seen_reference_keys.add(key)

    if changed:
        db.session.commit()


@app.context_processor
def inject_globals():
    active_period = Period.query.filter_by(is_active=True).order_by(Period.created_at.desc()).first()
    return {"active_period": active_period}


@app.route("/")
def index():
    active_period = Period.query.filter_by(is_active=True).order_by(Period.created_at.desc()).first()
    periods = Period.query.order_by(Period.created_at.desc()).all()
    period_counts = {
        pid: count
        for pid, count in db.session.query(Entry.period_id, func.count(Entry.id)).group_by(Entry.period_id).all()
    }
    reference_count = db.session.query(func.count(ReferenceEntry.id)).scalar() or 0
    entry_count = db.session.query(func.count(Entry.id)).scalar() or 0
    latest_entries = []
    if active_period:
        latest_entries = (
            Entry.query.filter_by(period_id=active_period.id)
            .order_by(Entry.created_at.desc())
            .limit(10)
            .all()
        )
    return render_template(
        "index.html",
        active_period=active_period,
        periods=periods,
        period_counts=period_counts,
        reference_count=reference_count,
        entry_count=entry_count,
        latest_entries=latest_entries,
    )





@app.route("/api/entries-search")
def api_entries_search():
    """Search the full Entry table across all Entry columns, optionally limited by period_id."""
    q = (request.args.get("q") or "").strip()
    period_id = request.args.get("period_id", type=int)
    limit = request.args.get("limit", 5000, type=int) or 5000
    limit = max(1, min(limit, 10000))

    query = Entry.query
    if period_id:
        query = query.filter(Entry.period_id == period_id)

    if q:
        pattern = f"%{q}%"
        filters = []
        for column in Entry.__table__.columns:
            try:
                filters.append(column.cast(db.String).ilike(pattern))
            except Exception:
                pass
        # Also allow searching by the month name connected to the receipt.
        try:
            query = query.join(Period, Entry.period_id == Period.id)
            filters.append(Period.name.ilike(pattern))
        except Exception:
            pass
        if filters:
            query = query.filter(or_(*filters))

    query = query.order_by(Entry.created_at.desc(), Entry.id.desc())

    total = query.count()
    rows = query.limit(limit).all()

    def value(row, name):
        val = getattr(row, name, "")
        return "" if val is None else val

    def money(row, name):
        val = getattr(row, name, None)
        if val is None:
            return ""
        try:
            return f"{float(val):,.2f}"
        except Exception:
            return str(val)

    return jsonify({
        "items": [
            {
                "id": value(row, "id"),
                "period_id": value(row, "period_id"),
                "period_name": row.period.name if getattr(row, "period", None) else "",
                "category": value(row, "category"),
                "store": value(row, "store"),
                "address": value(row, "address"),
                "actual_tin_number": value(row, "actual_tin_number"),
                "tin_number_on_input_tax": value(row, "tin_number_on_input_tax"),
                "invoice_value": money(row, "invoice_value"),
                "vat_12": money(row, "vat_12"),
                "cost_wo_vat": money(row, "cost_wo_vat"),
                "created_at": row.created_at.strftime("%Y-%m-%d %H:%M:%S") if row.created_at else "",
                "updated_at": row.updated_at.strftime("%Y-%m-%d %H:%M:%S") if row.updated_at else "",
            }
            for row in rows
        ],
        "total": total,
        "limit": limit,
    })


@app.route("/api/reference-summary-search")
def api_reference_summary_search():
    """Search the full ReferenceEntry table across all ReferenceEntry columns."""
    q = (request.args.get("q") or "").strip()
    limit = request.args.get("limit", 5000, type=int) or 5000
    limit = max(1, min(limit, 10000))

    query = ReferenceEntry.query
    if q:
        pattern = f"%{q}%"
        filters = []
        for column in ReferenceEntry.__table__.columns:
            try:
                filters.append(column.cast(db.String).ilike(pattern))
            except Exception:
                pass
        if filters:
            query = query.filter(or_(*filters))

    query = query.order_by(ReferenceEntry.uploaded_at.desc(), ReferenceEntry.id.desc())

    total = query.count()
    rows = query.limit(limit).all()

    def value(row, name):
        val = getattr(row, name, "")
        return "" if val is None else val

    return jsonify({
        "items": [
            {
                "id": value(row, "id"),
                "actual_tin_receipt": value(row, "actual_tin_receipt"),
                "actual_tin_normalized": value(row, "actual_tin_normalized"),
                "tin_needed_input_tax": value(row, "tin_needed_input_tax"),
                "category": value(row, "category"),
                "store": value(row, "store"),
                "address": value(row, "address"),
                "uploaded_at": row.uploaded_at.strftime("%Y-%m-%d %H:%M:%S") if row.uploaded_at else "",
            }
            for row in rows
        ],
        "total": total,
        "limit": limit,
    })


@app.route("/periods", methods=["GET", "POST"])
def manage_periods():
    if request.method == "POST":
        action = request.form.get("action", "create")
        if action == "create":
            name = (request.form.get("name") or "").strip()
            if not name:
                flash("Please enter a month name like Mar'26.", "danger")
                return redirect(url_for("manage_periods"))
            existing = Period.query.filter(func.lower(Period.name) == name.lower()).first()
            if existing:
                flash(f"Period {name} already exists.", "warning")
                return redirect(url_for("manage_periods"))
            period = Period(name=name, is_active=bool(request.form.get("set_active")))
            if period.is_active:
                Period.query.update({Period.is_active: False})
            db.session.add(period)
            db.session.commit()
            flash(f"Period {name} created.", "success")
            return redirect(url_for("manage_periods"))

        if action == "activate":
            period_id = request.form.get("period_id", type=int)
            period = db.session.get(Period, period_id)
            if not period:
                flash("Selected month does not exist.", "danger")
                return redirect(url_for("manage_periods"))
            Period.query.update({Period.is_active: False})
            period.is_active = True
            db.session.commit()
            flash(f"{period.name} is now the active month.", "success")
            return redirect(url_for("manage_periods"))

    periods = Period.query.order_by(Period.created_at.desc()).all()
    period_counts = {
        pid: count
        for pid, count in db.session.query(Entry.period_id, func.count(Entry.id)).group_by(Entry.period_id).all()
    }
    return render_template("periods.html", periods=periods, period_counts=period_counts)


@app.route("/reference", methods=["GET", "POST"])
def upload_reference():
    if request.method == "POST":
        action = request.form.get("action", "upload")

        if action == "upload":
            uploaded = request.files.get("reference_file")
            replace_all = bool(request.form.get("replace_all"))
            if not uploaded or not uploaded.filename:
                flash("Please choose an Excel file to upload.", "danger")
                return redirect(url_for("upload_reference"))

            suffix = Path(uploaded.filename).suffix.lower()
            if suffix not in {".xlsx", ".xlsm"}:
                flash("Please upload an .xlsx or .xlsm file.", "danger")
                return redirect(url_for("upload_reference"))

            saved_path = UPLOAD_DIR / f"reference_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}{suffix}"
            uploaded.save(saved_path)

            try:
                imported = import_reference_workbook(saved_path, replace_all=replace_all)
            except Exception as exc:
                flash(f"Import failed: {exc}", "danger")
                return redirect(url_for("upload_reference"))

            flash(f"Reference file imported successfully. {imported} rows loaded.", "success")
            return redirect(url_for("upload_reference"))

        if action == "manual_add":
            error = save_reference_from_form(request.form)
            if error:
                flash(error, "danger")
            else:
                flash("Reference row saved.", "success")
            return redirect(url_for("upload_reference"))

    count = db.session.query(func.count(ReferenceEntry.id)).scalar() or 0
    categories = get_reference_categories()
    page = request.args.get("page", 1, type=int)
    latest_pagination = ReferenceEntry.query.order_by(ReferenceEntry.uploaded_at.desc(), ReferenceEntry.id.desc()).paginate(page=page, per_page=15, error_out=False)
    latest = latest_pagination.items
    return render_template("reference.html", count=count, latest=latest, categories=categories, edit_row=None, latest_pagination=latest_pagination)


@app.route("/reference/<int:reference_id>/edit", methods=["GET", "POST"])
def edit_reference(reference_id: int):
    row = db.session.get(ReferenceEntry, reference_id)
    if not row:
        flash("Reference row not found.", "danger")
        return redirect(url_for("upload_reference"))

    if request.method == "POST":
        error = save_reference_from_form(request.form, existing=row)
        if error:
            flash(error, "danger")
            categories = get_reference_categories()
            count = db.session.query(func.count(ReferenceEntry.id)).scalar() or 0
            latest_pagination = ReferenceEntry.query.order_by(ReferenceEntry.uploaded_at.desc(), ReferenceEntry.id.desc()).paginate(page=1, per_page=15, error_out=False)
            latest = latest_pagination.items
            return render_template("reference.html", count=count, latest=latest, categories=categories, edit_row=row, latest_pagination=latest_pagination)
        flash("Reference row updated.", "success")
        return redirect(url_for("upload_reference"))

    count = db.session.query(func.count(ReferenceEntry.id)).scalar() or 0
    categories = get_reference_categories()
    latest_pagination = ReferenceEntry.query.order_by(ReferenceEntry.uploaded_at.desc(), ReferenceEntry.id.desc()).paginate(page=1, per_page=15, error_out=False)
    latest = latest_pagination.items
    return render_template("reference.html", count=count, latest=latest, categories=categories, edit_row=row, latest_pagination=latest_pagination)


@app.post("/reference/<int:reference_id>/delete")
def delete_reference(reference_id: int):
    row = db.session.get(ReferenceEntry, reference_id)
    if not row:
        flash("Reference row not found.", "danger")
        return redirect(url_for("upload_reference"))
    db.session.delete(row)
    db.session.commit()
    flash("Reference row deleted.", "success")
    return redirect(url_for("upload_reference"))


@app.route("/entries")
def list_entries():
    period_id = request.args.get("period_id", type=int)
    periods = Period.query.order_by(Period.created_at.desc()).all()
    if not period_id:
        active = Period.query.filter_by(is_active=True).first()
        if active:
            period_id = active.id
        elif periods:
            period_id = periods[0].id

    selected_period = db.session.get(Period, period_id) if period_id else None
    page = request.args.get("page", 1, type=int)
    entries = []
    entries_pagination = None
    if selected_period:
        entries_pagination = Entry.query.filter_by(period_id=selected_period.id).order_by(Entry.created_at.desc(), Entry.id.desc()).paginate(page=page, per_page=15, error_out=False)
        entries = entries_pagination.items
    return render_template("entries.html", periods=periods, selected_period=selected_period, entries=entries, entries_pagination=entries_pagination)


@app.route("/entries/new", methods=["GET", "POST"])
def new_entry():
    periods = Period.query.order_by(Period.created_at.desc()).all()
    default_period = Period.query.filter_by(is_active=True).first() or (periods[0] if periods else None)

    if request.method == "POST":
        period_id = request.form.get("period_id", type=int)
        period = db.session.get(Period, period_id) if period_id else default_period
        if not period:
            flash("Please create a month first.", "danger")
            return redirect(url_for("manage_periods"))

        actual_tin = (request.form.get("actual_tin_number") or "").strip()
        if not actual_tin:
            flash("Actual TIN Number is required.", "danger")
            return redirect(url_for("new_entry"))

        try:
            invoice_value = to_decimal(request.form.get("invoice_value"))
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("new_entry"))

        reference = find_reference(actual_tin)
        vat, cost = compute_vat_and_cost(invoice_value)

        entry = Entry(
            period_id=period.id,
            category=reference.category if reference else None,
            store=reference.store if reference else None,
            address=reference.address if reference else None,
            actual_tin_number=actual_tin,
            tin_number_on_input_tax=reference.tin_needed_input_tax if reference else None,
            invoice_value=invoice_value,
            vat_12=vat,
            cost_wo_vat=cost,
        )
        db.session.add(entry)
        db.session.commit()
        if reference:
            flash("Entry receipt saved and matched to the reference file.", "success")
        else:
            flash("Entry receipt saved, but no reference match was found for that TIN.", "warning")
        return redirect(url_for("total_input_vat"))

    return render_template("entry_form.html", periods=periods, entry=None, default_period=default_period)


@app.route("/entries/<int:entry_id>/edit", methods=["GET", "POST"])
def edit_entry(entry_id: int):
    entry = db.session.get(Entry, entry_id)
    if not entry:
        flash("Entry receipt not found.", "danger")
        return redirect(url_for("total_input_vat"))
    periods = Period.query.order_by(Period.created_at.desc()).all()

    if request.method == "POST":
        period_id = request.form.get("period_id", type=int)
        period = db.session.get(Period, period_id)
        if not period:
            flash("Selected month does not exist.", "danger")
            return redirect(url_for("edit_entry", entry_id=entry.id))

        actual_tin = (request.form.get("actual_tin_number") or "").strip()
        if not actual_tin:
            flash("Actual TIN Number is required.", "danger")
            return redirect(url_for("edit_entry", entry_id=entry.id))

        try:
            invoice_value = to_decimal(request.form.get("invoice_value"))
        except ValueError as exc:
            flash(str(exc), "danger")
            return redirect(url_for("edit_entry", entry_id=entry.id))

        reference = find_reference(actual_tin)
        vat, cost = compute_vat_and_cost(invoice_value)

        entry.period_id = period.id
        entry.actual_tin_number = actual_tin
        entry.invoice_value = invoice_value
        entry.vat_12 = vat
        entry.cost_wo_vat = cost
        entry.category = reference.category if reference else None
        entry.store = reference.store if reference else None
        entry.address = reference.address if reference else None
        entry.tin_number_on_input_tax = reference.tin_needed_input_tax if reference else None

        db.session.commit()
        if reference:
            flash("Entry receipt updated and matched to the reference file.", "success")
        else:
            flash("Entry receipt updated, but no reference match was found for that TIN.", "warning")
        return redirect(url_for("total_input_vat"))

    return render_template("entry_form.html", periods=periods, entry=entry, default_period=entry.period)


@app.post("/entries/<int:entry_id>/delete")
def delete_entry(entry_id: int):
    entry = db.session.get(Entry, entry_id)
    if not entry:
        flash("Entry receipt not found.", "danger")
        return redirect(url_for("total_input_vat"))
    period_id = entry.period_id
    db.session.delete(entry)
    db.session.commit()
    flash("Entry receipt deleted.", "success")
    return redirect(url_for("total_input_vat"))


@app.get("/api/lookup")
def api_lookup():
    tin = request.args.get("tin", "")
    reference = find_reference(tin)
    invoice = request.args.get("invoice", "")
    vat = cost = None
    if invoice not in {None, ""}:
        try:
            invoice_dec = to_decimal(invoice)
            vat_dec, cost_dec = compute_vat_and_cost(invoice_dec)
            vat = f"{vat_dec:.2f}"
            cost = f"{cost_dec:.2f}"
        except ValueError:
            pass

    if not reference:
        return jsonify({"matched": False, "vat": vat, "cost_wo_vat": cost})

    return jsonify(
        {
            "matched": True,
            "category": reference.category,
            "store": reference.store,
            "address": reference.address,
            "tin_number_on_input_tax": reference.tin_needed_input_tax,
            "vat": vat,
            "cost_wo_vat": cost,
        }
    )


@app.get("/export/<int:period_id>")
def export_period(period_id: int):
    period = db.session.get(Period, period_id)
    if not period:
        flash("Selected month does not exist.", "danger")
        return redirect(url_for("total_input_vat"))
    workbook = build_period_workbook(period)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"{period.name.replace('/', '-')}_export.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def find_reference(actual_tin: str | None) -> Optional[ReferenceEntry]:
    normalized = normalize_tin(actual_tin)
    if not normalized:
        return None
    reference = ReferenceEntry.query.filter_by(actual_tin_normalized=normalized).order_by(ReferenceEntry.id.asc()).first()
    if reference:
        return reference
    return ReferenceEntry.query.filter_by(actual_tin_receipt=(actual_tin or "").strip()).order_by(ReferenceEntry.id.asc()).first()


def save_reference_from_form(form, existing: ReferenceEntry | None = None) -> Optional[str]:
    actual_tin = (form.get("actual_tin_receipt") or "").strip()
    tin_needed = (form.get("tin_needed_input_tax") or "").strip()
    store = (form.get("store") or "").strip()
    address = (form.get("address") or "").strip()
    category = (form.get("category") or "").strip()
    custom_category = (form.get("custom_category") or "").strip()

    if category == "__custom__":
        category = custom_category
    elif not category and custom_category:
        category = custom_category
    category = normalize_category(category)

    normalized = normalize_tin(actual_tin)
    if not all([actual_tin, tin_needed, store, address, category, normalized]):
        return "Please complete Actual TIN, TIN Needed, Store, Address, and CAT."

    duplicate = ReferenceEntry.query.filter_by(actual_tin_normalized=normalized).first()
    if duplicate and (existing is None or duplicate.id != existing.id):
        return "That Actual TIN already exists in reference. Edit the existing row instead."

    row = existing or ReferenceEntry()
    row.actual_tin_receipt = actual_tin
    row.actual_tin_normalized = normalized
    row.tin_needed_input_tax = tin_needed
    row.store = store
    row.address = address
    row.category = category
    row.uploaded_at = datetime.utcnow()

    if existing is None:
        db.session.add(row)

    db.session.commit()
    return None


def import_reference_workbook(path: Path, replace_all: bool = True) -> int:
    wb = load_workbook(path, data_only=True)

    ws = wb["REFERENCE"] if "REFERENCE" in wb.sheetnames else None
    if ws is None:
        for candidate in wb.worksheets:
            headers = [canonical_header(cell.value) for cell in candidate[1]]
            if REQUIRED_HEADER_KEYS.issubset(set(headers)):
                ws = candidate
                break

    if ws is None:
        raise ValueError("Could not find a REFERENCE sheet with the required columns.")

    raw_headers = [cell.value for cell in ws[1]]
    headers = [canonical_header(cell) for cell in raw_headers]
    target_headers = {name: idx for idx, name in enumerate(headers) if name}

    if not REQUIRED_HEADER_KEYS.issubset(target_headers.keys()):
        raise ValueError(
            "The reference file must contain ACTUAL TIN NUMBER ON RECEIPT, TIN NUMBER NEEDED ON INPUT TAX, STORE, ADDRESS, and CAT columns."
        )

    rows_to_insert: list[ReferenceEntry] = []
    seen: set[str] = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        actual_tin = str(row[target_headers["actual_tin_number_on_receipt"]] or "").strip()
        tin_needed = str(row[target_headers["tin_number_needed_on_input_tax"]] or "").strip()
        store = str(row[target_headers["store"]] or "").strip()
        address = str(row[target_headers["address"]] or "").strip()
        category = normalize_category(str(row[target_headers["cat"]] or "").strip())

        normalized = normalize_tin(actual_tin)
        if not all([actual_tin, tin_needed, store, address, category, normalized]):
            continue

        dedupe_key = f"{normalized}|{tin_needed}|{store}|{address}|{category}"
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)

        rows_to_insert.append(
            ReferenceEntry(
                actual_tin_receipt=actual_tin,
                actual_tin_normalized=normalized,
                tin_needed_input_tax=tin_needed,
                store=store,
                address=address,
                category=category,
            )
        )

    if replace_all:
        db.session.query(ReferenceEntry).delete()

    db.session.add_all(rows_to_insert)
    db.session.commit()
    return len(rows_to_insert)


def build_period_workbook(period: Period) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = period.name

    headers = [
        "aaaa",
        "CAT",
        "STORE",
        "ADDRESS",
        "ACTUAL TIN NUMBER",
        "TIN NUMBER ON INPUT TAX",
        "INVOICE VALUE",
        "12%VAT",
        "COST W/O VAT",
    ]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)

    entries = Entry.query.filter_by(period_id=period.id).order_by(Entry.id.asc()).all()
    currency_format = '#,##0.00_);[Red](#,##0.00)'

    for row_idx, entry in enumerate(entries, start=2):
        ws.cell(row=row_idx, column=1, value="")
        ws.cell(row=row_idx, column=2, value=entry.category)
        ws.cell(row=row_idx, column=3, value=entry.store)
        ws.cell(row=row_idx, column=4, value=entry.address)
        ws.cell(row=row_idx, column=5, value=entry.actual_tin_number)
        ws.cell(row=row_idx, column=6, value=entry.tin_number_on_input_tax)
        invoice_value = to_decimal(entry.invoice_value)
        vat_value = to_decimal(entry.vat_12) if entry.vat_12 is not None else Decimal("0.00")
        cost_value = to_decimal(entry.cost_wo_vat) if entry.cost_wo_vat is not None else Decimal("0.00")
        if vat_value == Decimal("0.00") and cost_value == Decimal("0.00") and invoice_value != Decimal("0.00"):
            vat_value, cost_value = compute_vat_and_cost(invoice_value)

        ws.cell(row=row_idx, column=7, value=float(invoice_value))
        ws.cell(row=row_idx, column=8, value=float(vat_value))
        ws.cell(row=row_idx, column=9, value=float(cost_value))

        ws.cell(row=row_idx, column=7).number_format = currency_format
        ws.cell(row=row_idx, column=8).number_format = currency_format
        ws.cell(row=row_idx, column=9).number_format = currency_format

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(len(entries) + 1, 2)}"
    widths = {
        "A": 12,
        "B": 10,
        "C": 34,
        "D": 42,
        "E": 24,
        "F": 24,
        "G": 16,
        "H": 14,
        "I": 16,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    return wb


with app.app_context():
    db.create_all()
    cleanup_existing_categories()



@app.route("/reference/export")
def export_references_excel():
    """Export the full Reference Summary database to Excel."""
    references = ReferenceEntry.query.order_by(ReferenceEntry.id.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Reference Summary"

    headers = [
        "Actual TIN Number on Receipt",
        "TIN Number Needed on Input Tax",
        "CAT",
        "Store",
        "Address",
    ]
    ws.append(headers)

    for ref in references:
        ws.append([
            getattr(ref, "actual_tin_receipt", "") or "",
            getattr(ref, "tin_needed_input_tax", "") or "",
            getattr(ref, "category", "") or "",
            getattr(ref, "store", "") or "",
            getattr(ref, "address", "") or "",
        ])

    for column_cells in ws.columns:
        column_letter = column_cells[0].column_letter
        max_length = 0
        for cell in column_cells:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max_length + 3, 70)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="reference_summary.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# --- Total Input of VAT page (active month only) ---
def get_active_period() -> Optional[Period]:
    """Return the currently active month/period."""
    return Period.query.filter_by(is_active=True).order_by(Period.created_at.desc()).first()


def apply_entry_full_column_search(query, q: str):
    """Search every Entry database column plus the connected month name."""
    q = (q or "").strip()
    if not q:
        return query
    pattern = f"%{q}%"
    filters = []
    for column in Entry.__table__.columns:
        try:
            filters.append(column.cast(db.String).ilike(pattern))
        except Exception:
            pass
    try:
        query = query.join(Period, Entry.period_id == Period.id)
        filters.append(Period.name.ilike(pattern))
    except Exception:
        pass
    if filters:
        query = query.filter(or_(*filters))
    return query


def total_input_vat_entries_query(period: Period, q: str = ""):
    """Query for Total Input of VAT rows, active month only, sorted A to Z by Store."""
    query = Entry.query.filter(Entry.period_id == period.id)
    query = apply_entry_full_column_search(query, q)
    return query.order_by(
        func.lower(func.coalesce(Entry.store, "")).asc(),
        func.lower(func.coalesce(Entry.category, "")).asc(),
        Entry.id.asc(),
    )


def query_total_input_vat_entries(period: Period, q: str = "") -> list[Entry]:
    """Rows for the Total Input of VAT page/export, sorted A to Z by Store."""
    return total_input_vat_entries_query(period, q).all()


def total_input_vat_summary(period: Period) -> dict[str, Decimal | int]:
    """Live totals for the active month."""
    totals = (
        db.session.query(
            func.count(Entry.id),
            func.coalesce(func.sum(Entry.invoice_value), 0),
            func.coalesce(func.sum(Entry.vat_12), 0),
            func.coalesce(func.sum(Entry.cost_wo_vat), 0),
        )
        .filter(Entry.period_id == period.id)
        .first()
    )
    return {
        "entry_count": int(totals[0] or 0),
        "total_invoice": to_decimal(totals[1] or 0),
        "total_vat": to_decimal(totals[2] or 0),
        "total_cost": to_decimal(totals[3] or 0),
    }


def money_text(value) -> str:
    return f"{to_decimal(value):,.2f}"


def total_input_vat_item(entry: Entry) -> dict[str, str | int]:
    return {
        "id": entry.id,
        "period_id": entry.period_id,
        "period_name": entry.period.name if entry.period else "",
        "category": entry.category or "",
        "store": entry.store or "",
        "address": entry.address or "",
        "actual_tin_number": entry.actual_tin_number or "",
        "tin_number_on_input_tax": entry.tin_number_on_input_tax or "",
        "invoice_value": money_text(entry.invoice_value),
        "vat_12": money_text(entry.vat_12),
        "cost_wo_vat": money_text(entry.cost_wo_vat),
        "created_at": entry.created_at.strftime("%Y-%m-%d %H:%M:%S") if entry.created_at else "",
        "updated_at": entry.updated_at.strftime("%Y-%m-%d %H:%M:%S") if entry.updated_at else "",
    }


@app.route("/total-input-vat")
def total_input_vat():
    """Show active month input VAT records with pagination, full-column search, and live VAT total."""
    active = get_active_period()
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 15, type=int) or 15
    per_page = max(5, min(per_page, 100))
    q = (request.args.get("q") or "").strip()

    entries = []
    total_pagination = None
    filtered_count = 0
    summary = {
        "entry_count": 0,
        "total_invoice": Decimal("0.00"),
        "total_vat": Decimal("0.00"),
        "total_cost": Decimal("0.00"),
    }

    if active:
        summary = total_input_vat_summary(active)
        total_pagination = total_input_vat_entries_query(active, q).paginate(page=page, per_page=per_page, error_out=False)
        entries = total_pagination.items
        filtered_count = total_pagination.total

    return render_template(
        "total_input_vat.html",
        active_period=active,
        entries=entries,
        summary=summary,
        filtered_count=filtered_count,
        total_pagination=total_pagination,
        q=q,
        per_page=per_page,
        money_text=money_text,
    )


@app.get("/api/total-input-vat")
def api_total_input_vat():
    """JSON used by the Total Input VAT page for auto-refresh, search, and pagination."""
    active = get_active_period()
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 15, type=int) or 15
    per_page = max(5, min(per_page, 100))
    q = (request.args.get("q") or "").strip()

    if not active:
        return jsonify({
            "active_period": None,
            "summary": {
                "entry_count": 0,
                "total_invoice": "0.00",
                "total_vat": "0.00",
                "total_cost": "0.00",
            },
            "pagination": {"page": 1, "pages": 1, "per_page": per_page, "total": 0, "has_prev": False, "has_next": False},
            "items": [],
        })

    pagination = total_input_vat_entries_query(active, q).paginate(page=page, per_page=per_page, error_out=False)
    summary = total_input_vat_summary(active)
    return jsonify({
        "active_period": {"id": active.id, "name": active.name},
        "summary": {
            "entry_count": summary["entry_count"],
            "total_invoice": money_text(summary["total_invoice"]),
            "total_vat": money_text(summary["total_vat"]),
            "total_cost": money_text(summary["total_cost"]),
        },
        "pagination": {
            "page": pagination.page,
            "pages": pagination.pages or 1,
            "per_page": pagination.per_page,
            "total": pagination.total,
            "has_prev": pagination.has_prev,
            "has_next": pagination.has_next,
            "prev_num": pagination.prev_num if pagination.has_prev else None,
            "next_num": pagination.next_num if pagination.has_next else None,
        },
        "items": [total_input_vat_item(entry) for entry in pagination.items],
    })


@app.get("/total-input-vat/export")
def export_total_input_vat():
    """Export active month Total Input of VAT records to Excel, sorted A to Z."""
    active = get_active_period()
    if not active:
        flash("Please activate a month first before exporting Total Input of VAT.", "warning")
        return redirect(url_for("total_input_vat"))

    wb = build_total_input_vat_workbook(active)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", active.name).strip("_") or "active_month"
    return send_file(
        output,
        as_attachment=True,
        download_name=f"total_input_vat_{safe_name}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def build_total_input_vat_workbook(period: Period) -> Workbook:
    """Excel version of the Total Input of VAT page."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Total Input VAT"

    headers = [
        "CAT",
        "STORE",
        "ADDRESS",
        "ACTUAL TIN NUMBER",
        "TIN NUMBER ON INPUT TAX",
        "INVOICE VALUE",
        "12%VAT",
        "COST W/O VAT",
    ]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    total_font = Font(bold=True)
    thin = Side(style="thin", color="D9E2F3")
    currency_format = '#,##0.00_);[Red](#,##0.00)'

    ws.cell(row=1, column=1, value="Total Input of VAT")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Active Month")
    ws.cell(row=2, column=2, value=period.name)

    header_row = 4
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)

    entries = query_total_input_vat_entries(period)
    start_row = header_row + 1
    for row_idx, entry in enumerate(entries, start=start_row):
        invoice_value = to_decimal(entry.invoice_value)
        vat_value = to_decimal(entry.vat_12) if entry.vat_12 is not None else Decimal("0.00")
        cost_value = to_decimal(entry.cost_wo_vat) if entry.cost_wo_vat is not None else Decimal("0.00")

        ws.cell(row=row_idx, column=1, value=entry.category)
        ws.cell(row=row_idx, column=2, value=entry.store)
        ws.cell(row=row_idx, column=3, value=entry.address)
        ws.cell(row=row_idx, column=4, value=entry.actual_tin_number)
        ws.cell(row=row_idx, column=5, value=entry.tin_number_on_input_tax)
        ws.cell(row=row_idx, column=6, value=float(invoice_value))
        ws.cell(row=row_idx, column=7, value=float(vat_value))
        ws.cell(row=row_idx, column=8, value=float(cost_value))
        for col_idx in (6, 7, 8):
            ws.cell(row=row_idx, column=col_idx).number_format = currency_format

    total_row = start_row + len(entries)
    summary = total_input_vat_summary(period)
    ws.cell(row=total_row, column=5, value="TOTAL")
    ws.cell(row=total_row, column=6, value=float(summary["total_invoice"]))
    ws.cell(row=total_row, column=7, value=float(summary["total_vat"]))
    ws.cell(row=total_row, column=8, value=float(summary["total_cost"]))
    for col_idx in range(5, 9):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.fill = total_fill
        cell.font = total_font
        cell.border = Border(top=thin, bottom=thin)
        if col_idx >= 6:
            cell.number_format = currency_format

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:H{max(total_row, header_row + 1)}"
    widths = {
        "A": 12,
        "B": 34,
        "C": 42,
        "D": 24,
        "E": 24,
        "F": 16,
        "G": 14,
        "H": 16,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    return wb


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)

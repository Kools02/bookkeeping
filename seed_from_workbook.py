from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app import app, db, import_reference_workbook, Period, Entry, compute_vat_and_cost

WORKBOOK = Path(r"/mnt/data/March'26_step1.xlsx")
SHEET_NAME = "Mar'26"

with app.app_context():
    import_reference_workbook(WORKBOOK, replace_all=True)

    period = Period.query.filter_by(name=SHEET_NAME).first()
    if not period:
        Period.query.update({Period.is_active: False})
        period = Period(name=SHEET_NAME, is_active=True)
        db.session.add(period)
        db.session.commit()
    else:
        Period.query.update({Period.is_active: False})
        period.is_active = True
        db.session.commit()

    Entry.query.filter_by(period_id=period.id).delete()
    db.session.commit()

    wb = load_workbook(WORKBOOK, data_only=True)
    ws = wb[SHEET_NAME]
    inserted = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        actual_tin = row[5]
        invoice_value = row[7]
        if not actual_tin or invoice_value in (None, ""):
            continue

        try:
            invoice = Decimal(str(invoice_value))
        except Exception:
            continue

        vat, cost = compute_vat_and_cost(invoice)
        entry = Entry(
            period_id=period.id,
            date_coded=row[1],
            actual_date_on_receipt=None,
            category=row[2],
            store=row[3],
            address=row[4],
            actual_tin_number=str(actual_tin),
            tin_number_on_input_tax=row[6],
            invoice_value=invoice,
            vat_12=vat,
            cost_wo_vat=cost,
        )
        db.session.add(entry)
        inserted += 1

    db.session.commit()
    print(f"Seeded {inserted} entries into {SHEET_NAME}.")
